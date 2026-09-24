import docx

def md_to_docx(md_path, docx_path):
    doc = docx.Document()
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_table = False
    table_rows = []
    
    def render_table():
        if not table_rows: return
        cols = len(table_rows[0])
        table = doc.add_table(rows=len(table_rows), cols=cols)
        table.style = 'Table Grid'
        
        for i, row_data in enumerate(table_rows):
            row_cells = table.rows[i].cells
            for j, cell_data in enumerate(row_data):
                if j < len(row_cells):
                    text = cell_data.replace('**', '').strip()
                    row_cells[j].text = text
                    
        doc.add_paragraph("")
        table_rows.clear()
        
    for line in lines:
        line = line.strip()
        if not line:
            if in_table:
                render_table()
                in_table = False
            continue
            
        if line.startswith('---'):
            continue
            
        if line.startswith('|'):
            if '| :---' in line or '| ---' in line:
                continue
            in_table = True
            parts = [p.strip() for p in line.split('|')][1:-1]
            if parts:
                table_rows.append(parts)
        else:
            if in_table:
                render_table()
                in_table = False
                
            if line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            else:
                doc.add_paragraph(line.replace('**', ''))
                
    if in_table:
        render_table()
        
    doc.save(docx_path)

if __name__ == '__main__':
    md = r"C:\Users\miure\.gemini\antigravity-ide\brain\8b21a557-f969-4b90-9029-c5d149a901d1\Requerimientos_Panaderia_Amada.md"
    out = r"C:\Users\miure\.gemini\antigravity-ide\brain\8b21a557-f969-4b90-9029-c5d149a901d1\Requerimientos_Panaderia_Amada.docx"
    md_to_docx(md, out)
    print("DONE")
