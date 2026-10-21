"""Flujo DGI real sobre SQL Server aislado. Sin credenciales ni datos de producción."""
import hashlib
import io
import json
import os
import re
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
import calendar

import pyodbc
import pytest

from app import create_app
from app.models import User
from app.services import dgi_repository as repo
from config import Config
from scripts.demo_dgi import correct_case, require_qa, seed_case, create_demo_user
from scripts.migrate_dgi import downgrade, migrate


@pytest.fixture(scope='module')
def qa_string():
    value = os.environ.get('TEST_DB_CONNECTION_STRING')
    if not value:
        pytest.skip('Defina TEST_DB_CONNECTION_STRING para una base PanaderiaQA_ aislada.')
    connection = pyodbc.connect(value)
    try:
        require_qa(connection)
        migrate(connection)
        connection.commit()
    finally:
        connection.rollback()
        connection.autocommit = True
        connection.execute('SET TRANSACTION ISOLATION LEVEL READ COMMITTED')
        connection.execute('USE master')
        connection.close()
    return value


@pytest.fixture
def app(qa_string):
    class TestConfig(Config):
        TESTING = True
        SECRET_KEY = 'dgi-isolated-qa-only'
        SQL_SERVER_CONNECTION_STRING = qa_string
    application = create_app(TestConfig)
    application.login_manager._user_callback = lambda _: User(1, 'qa', 'Revisor sintético QA', 1, 'SuperAdmin')
    return application


@pytest.fixture
def client(app):
    result = app.test_client()
    with result.session_transaction() as session:
        session['_user_id'] = 'emp_1'
        session['_fresh'] = True
    return result


@pytest.fixture
def case(qa_string):
    connection = pyodbc.connect(qa_string)
    old_profile = repo.profile(connection.cursor())
    # Mes independiente por caso y ejecución. No toca operaciones de otras pruebas.
    for year in range(2040, 2100):
        for month in range(1, 13):
            stamp = date(year, month, 1)
            exists = connection.cursor().execute('''SELECT 1 WHERE
                EXISTS(SELECT 1 FROM DGI_Periodos WHERE Desde=?) OR
                EXISTS(SELECT 1 FROM Facturas WHERE FechaHora>=? AND FechaHora<DATEADD(month,1,?)) OR
                EXISTS(SELECT 1 FROM AsientosDiario WHERE Fecha>=? AND Fecha<DATEADD(month,1,?))''',
                stamp,stamp,stamp,stamp,stamp).fetchone()
            if not exists:
                selected = f'{year}-{month:02d}'
                break
        else:
            continue
        break
    result = seed_case(connection, selected)
    connection.commit()
    result['conn'] = connection
    yield result
    connection.rollback()
    repo.save_profile(connection.cursor(), old_profile, 1)
    connection.commit()
    connection.close()


def csrf(client):
    response = client.get('/admin/dgi')
    assert response.status_code == 200
    return re.search(r'name="csrf-token" content="([^"]+)"', response.text).group(1)


def post(client, path, token, **data):
    return client.post(path, data=dict(csrf_token=token, **data))


def correct_and_approve(client, case, token):
    rid = correct_case(case['conn'], case['periodo_id'])
    case['conn'].commit()
    assert post(client, f'/admin/dgi/revisiones/{rid}/revisar', token, nota='Verificados todos los soportes y el ajuste sintético.').status_code == 302
    approved = post(client, f'/admin/dgi/revisiones/{rid}/aprobar', token)
    assert approved.status_code == 302, approved.text
    return rid


def test_120000_110000_flow_blocks_approval_and_export(client, case):
    token = csrf(client)
    rid = case['revision_id']
    page = client.get(f'/admin/dgi/revisiones/{rid}')
    assert page.status_code == 200
    assert all(text in page.text for text in ['C$ 120,000.00', 'C$ 110,000.00', 'C$ 10,000.00', 'CONCILIACION'])
    assert post(client, f'/admin/dgi/revisiones/{rid}/aprobar', token).status_code == 409
    assert post(client, f'/admin/dgi/revisiones/{rid}/revisar', token, nota='Revisado, todavía existe diferencia.').status_code == 302
    assert post(client, f'/admin/dgi/revisiones/{rid}/aprobar', token).status_code == 409
    assert post(client, f'/admin/dgi/revisiones/{rid}/exportar', token).status_code == 409


def test_correction_revalidates_keeps_history_and_exports_readable_excel(client, case):
    from openpyxl import load_workbook
    token = csrf(client)
    rid = correct_and_approve(client, case, token)
    page = client.get(f'/admin/dgi/revisiones/{rid}')
    assert 'C$ 0.00' in page.text
    result = post(client, f'/admin/dgi/revisiones/{rid}/exportar', token)
    assert result.status_code == 302
    download = client.get(result.location)
    assert download.status_code == 200
    book = load_workbook(io.BytesIO(download.data), data_only=True)
    assert book.sheetnames == ['Resumen','Ventas','Compras','Hallazgos','Trazabilidad','Supuestos']
    cursor = case['conn'].cursor()
    row = cursor.execute('SELECT SHA256,Archivo FROM DGI_Exportaciones WHERE RevisionID=?', rid).fetchone()
    assert hashlib.sha256(download.data).hexdigest() == row.SHA256
    assert download.data == bytes(row.Archivo)
    assert cursor.execute('SELECT Estado FROM DGI_Revisiones WHERE ID=?',case['revision_id']).fetchone()[0] == 'OBSOLETO'
    assert cursor.execute("SELECT COUNT(*) FROM BitacoraFiscal WHERE Entidad='DGI_Revision' AND EntidadID=?",rid).fetchone()[0] >= 5
    again = post(client, f'/admin/dgi/revisiones/{rid}/exportar', token)
    assert again.location == result.location
    assert client.get(again.location).data == download.data


def test_changed_source_invalidates_approval_and_old_download(client, case):
    token = csrf(client)
    rid = correct_and_approve(client, case, token)
    export = post(client, f'/admin/dgi/revisiones/{rid}/exportar', token)
    cursor = case['conn'].cursor()
    cursor.execute('UPDATE ComprasMateriaPrima SET CostoTotal=CostoTotal+1 WHERE ID=?',case['compra_id'])
    case['conn'].commit()
    assert 'Las fuentes de esta revisión cambiaron.' in client.get(f'/admin/dgi/revisiones/{rid}').text
    response = client.get(export.location)
    assert response.status_code == 409 and 'caducó' in response.text
    assert cursor.execute('SELECT Estado FROM DGI_Revisiones WHERE ID=?',rid).fetchone()[0] == 'OBSOLETO'
    assert post(client, f'/admin/dgi/revisiones/{rid}/aprobar', token).status_code == 409


def test_changed_profile_invalidates_before_approval(client, case):
    token = csrf(client)
    rid = correct_case(case['conn'], case['periodo_id'])
    case['conn'].commit()
    assert post(client, f'/admin/dgi/revisiones/{rid}/revisar', token, nota='Documentos revisados y conciliados.').status_code == 302
    case['conn'].cursor().execute("UPDATE ConfiguracionSistema SET Valor='Otro responsable QA' WHERE Clave='dgi_responsable'")
    case['conn'].commit()
    assert post(client, f'/admin/dgi/revisiones/{rid}/aprobar', token).status_code == 409
    assert case['conn'].cursor().execute('SELECT Estado FROM DGI_Revisiones WHERE ID=?',rid).fetchone()[0] == 'OBSOLETO'


def test_human_review_required_even_when_reconciled(client, case):
    token = csrf(client)
    rid = correct_case(case['conn'], case['periodo_id'])
    case['conn'].commit()
    assert post(client, f'/admin/dgi/revisiones/{rid}/aprobar', token).status_code == 409


def test_comment_records_actor_and_cannot_override_blocker(client, case):
    token = csrf(client)
    cursor = case['conn'].cursor()
    result = json.loads(cursor.execute('SELECT ResultadoJSON FROM DGI_Revisiones WHERE ID=?',case['revision_id']).fetchone()[0])
    finding = next(f for f in result['hallazgos'] if f['codigo'] == 'CONCILIACION')
    path = f'/admin/dgi/revisiones/{case["revision_id"]}'
    assert post(client, path+'/observaciones', token, hallazgo=finding['id'], comentario='Falta revisar un documento de ingreso.').status_code == 302
    assert post(client, path+'/observaciones', token, hallazgo='f'*24, comentario='Fuera de revisión').status_code == 404
    assert 'Falta revisar un documento' in client.get(path).text
    assert json.loads(cursor.execute('SELECT ResultadoJSON FROM DGI_Revisiones WHERE ID=?',case['revision_id']).fetchone()[0])['bloqueantes'] == 1


def test_period_close_requires_approval_and_reopen_preserves_data(client, case):
    token = csrf(client)
    path = f'/admin/dgi/periodos/{case["periodo_id"]}/estado'
    assert post(client, path, token, estado='Cerrado').status_code == 409
    rid = correct_and_approve(client, case, token)
    assert post(client, path, token, estado='Cerrado').status_code == 302
    assert post(client, '/admin/dgi/revisiones', token, periodo_id=case['periodo_id'], obligacion='Diagnóstico interno').status_code == 409
    assert post(client, path, token, estado='Abierto').status_code == 302
    assert client.get(f'/admin/dgi/revisiones/{rid}').status_code == 200


def test_official_export_remains_blocked_in_server(client, case):
    token = csrf(client)
    response = post(client, f'/admin/dgi/revisiones/{case["revision_id"]}/exportar-oficial',token)
    assert response.status_code == 409 and 'PENDIENTE DE VALIDACIÓN' in response.text


@pytest.mark.parametrize('role', ['Estandar','Operativo','Cliente','Invitado'])
def test_unauthorized_roles_cannot_read_or_write_dgi(app, client, role):
    token = csrf(client)
    app.login_manager._user_callback = lambda _: User(1,'qa','QA',2,role)
    assert client.get('/admin/dgi').status_code == 302
    for path in ['/admin/dgi/revisiones', '/admin/dgi/periodos', '/admin/dgi/perfil',
                 '/admin/dgi/revisiones/1/aprobar', '/admin/dgi/revisiones/1/exportar',
                 '/admin/dgi/revisiones/1/diagnostico.xlsx']:
        assert post(client,path,token).status_code == 302


def test_admin_can_validate_export_and_close_internal_review(app, client, case):
    token = csrf(client)
    app.login_manager._user_callback = lambda _: User(1,'qa','Preparador QA',2,'Admin')
    page = client.get(f'/admin/dgi/revisiones/{case["revision_id"]}')
    assert page.status_code == 200 and 'Admin y SuperAdmin pueden validar' in page.text
    assert post(client,f'/admin/dgi/revisiones/{case["revision_id"]}/aprobar',token).status_code == 409
    rid = correct_and_approve(client,case,token)
    exported = post(client,f'/admin/dgi/revisiones/{rid}/exportar',token)
    assert exported.status_code == 302 and client.get(exported.location).status_code == 200
    assert post(client,f'/admin/dgi/periodos/{case["periodo_id"]}/estado',token,estado='Cerrado').status_code == 302
    assert post(client,f'/admin/dgi/revisiones/{rid}/exportar-oficial',token).status_code == 302


def test_csrf_and_unknown_ids(client, case):
    assert client.post('/admin/dgi/periodos',data={'mes':'2035-01'}).status_code == 400
    assert client.get('/admin/dgi/revisiones/2147483647').status_code == 404
    token = csrf(client)
    assert post(client,'/admin/dgi/periodos',token,mes='bad').status_code == 400
    assert client.get('/admin/dgi?periodo=bad').status_code == 400


def test_current_source_links_load(client, case):
    for kind, rid in [('factura',case['facturas'][0]), ('compra',case['compra_id']),('asiento',case['asiento_id'])]:
        response = client.get(f'/admin/dgi/fuentes/{kind}/{rid}')
        assert response.status_code == 200 and 'Ir al módulo de origen' in response.text
    assert client.get(f'/admin/fiscal/facturas/{case["facturas"][0]}/imprimir').status_code == 200


def test_repeated_parallel_diagnosis_is_idempotent(app, case):
    def create(_):
        with app.app_context(), repo.transaction() as cursor:
            return repo.create_review(cursor,case['periodo_id'],'Diagnóstico interno',1)
    with ThreadPoolExecutor(max_workers=3) as pool:
        ids = list(pool.map(create,range(3)))
    assert set(ids) == {case['revision_id']}


def test_connection_returns_to_normal_isolation(app, qa_string):
    with app.app_context(), repo.transaction() as cursor:
        assert cursor.execute('SELECT transaction_isolation_level FROM sys.dm_exec_sessions WHERE session_id=@@SPID').fetchone()[0] == 4
    connection = pyodbc.connect(qa_string)
    try:
        assert connection.execute('SELECT transaction_isolation_level FROM sys.dm_exec_sessions WHERE session_id=@@SPID').fetchone()[0] == 2
        assert connection.execute('SELECT @@LOCK_TIMEOUT').fetchone()[0] == -1
    finally:
        connection.close()


def test_synthetic_account_logs_in_with_actual_database_role(app, qa_string, case):
    from app.models import load_user
    connection = pyodbc.connect(qa_string)
    access = create_demo_user(connection)
    connection.commit()
    connection.close()
    app.login_manager._user_callback = load_user
    client = app.test_client()
    login = client.get('/login')
    token = re.search(r'name="csrf_token"[^>]*value="([^"]+)"',login.text).group(1)
    response = post(client,'/login',token,identificador=access['usuario'],password=access['clave'])
    assert response.status_code == 302
    page = client.get(f'/admin/dgi/revisiones/{case["revision_id"]}')
    assert page.status_code == 200 and 'Validar revisión interna' in page.text
    assert 'Revisor DGI SINTÉTICO QA' in page.text


def test_profile_regime_change_does_not_leave_pos_without_required_emitter_data(client, case):
    token = csrf(client)
    cursor = case['conn'].cursor()
    data = repo.profile(cursor)
    previous = cursor.execute("SELECT Valor FROM ConfiguracionSistema WHERE Clave='fiscal_direccion'").fetchone()[0]
    try:
        cursor.execute("UPDATE ConfiguracionSistema SET Valor='' WHERE Clave='fiscal_direccion'")
        case['conn'].commit()
        data['fiscal_regimen'] = 'General'
        response = post(client,'/admin/dgi/perfil',token,**data)
        assert response.status_code == 409 and 'requisitos del POS' in response.text
        assert repo.profile(cursor)['fiscal_regimen'] == 'Cuota fija'
    finally:
        cursor.execute("UPDATE ConfiguracionSistema SET Valor=? WHERE Clave='fiscal_direccion'",previous)
        case['conn'].commit()


def test_diagnostic_download_includes_errors_without_approving_or_changing_state(client, case):
    from openpyxl import load_workbook
    token = csrf(client)
    rid = case['revision_id']
    response = post(client,f'/admin/dgi/revisiones/{rid}/diagnostico.xlsx',token)
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.data))
    assert 'DIAGNÓSTICO SIN APROBACIÓN' in book['Resumen']['A2'].value
    assert any(c.value == 'CONCILIACION' for row in book['Hallazgos'] for c in row)
    cursor = case['conn'].cursor()
    assert cursor.execute('SELECT Estado FROM DGI_Revisiones WHERE ID=?',rid).fetchone()[0] == 'BORRADOR'
    assert cursor.execute('SELECT COUNT(*) FROM DGI_Exportaciones WHERE RevisionID=?',rid).fetchone()[0] == 0
    event = cursor.execute("SELECT TOP 1 Detalle FROM BitacoraFiscal WHERE EntidadID=? AND Accion='DGI_REPORTE_DIAGNOSTICO' ORDER BY ID DESC",rid).fetchone()
    assert json.loads(event[0])['sha256'] == hashlib.sha256(response.data).hexdigest()
    assert post(client,f'/admin/dgi/revisiones/{rid}/aprobar',token).status_code == 409


def test_partial_profile_allows_internal_flow_for_admin_without_invented_tax_identity(app,client,case):
    token = csrf(client)
    correct_case(case['conn'],case['periodo_id'])
    case['conn'].commit()
    app.login_manager._user_callback = lambda _: User(1,'qa','Administrador QA',2,'Admin')
    partial = {key:'' for key in repo.PROFILE_KEYS}
    partial.update(fiscal_regimen='Pendiente',dgi_responsable='Revisor QA')
    assert post(client,'/admin/dgi/perfil',token,**partial).status_code == 302
    response = post(client,'/admin/dgi/revisiones',token,periodo_id=case['periodo_id'],obligacion='Diagnóstico interno')
    assert response.status_code == 302
    rid = int(response.location.rsplit('/',1)[1])
    page = client.get(response.location)
    assert 'Sin errores en las comprobaciones de los datos' in page.text
    assert 'Configuración pendiente' in page.text
    assert 'No se detectaron inconsistencias de datos' in page.text
    assert post(client,f'/admin/dgi/revisiones/{rid}/revisar',token,nota='Revisados documentos internos; identificación fiscal pendiente.').status_code == 302
    assert post(client,f'/admin/dgi/revisiones/{rid}/aprobar',token).status_code == 302
    assert post(client,f'/admin/dgi/revisiones/{rid}/exportar',token).status_code == 302
    assert repo.profile(case['conn'].cursor())['fiscal_ruc'] == ''


def test_legacy_utc_invoice_uses_confirmed_receipt_month_without_editing_source(case):
    from app.services.contabilidad_service import ContabilidadService as Accounting
    cursor = case['conn'].cursor()
    start, end = repo.monthly_period(case['mes'])
    local = datetime(end.year,end.month,end.day,20,13,18)
    original = local + timedelta(hours=6,seconds=1)
    number = 'LEG-QA-'+uuid.uuid4().hex[:10]
    product = cursor.execute('SELECT TOP 1 ID FROM Productos ORDER BY ID').fetchone()[0]
    entry = Accounting.crear_asiento(cursor,'QA histórico UTC',[(Accounting.CAJA_GENERAL_CS,7,0),(Accounting.VENTAS_PAN,0,7)],1,fecha=local,referencia_externa=number)
    rid = cursor.execute('''INSERT INTO Facturas(NumeroFactura,CodigoSeguimiento,FechaHora,UsuarioID,Subtotal,IVA,Total)
        OUTPUT INSERTED.ID VALUES(?,?,?,1,7,0,7)''',number,uuid.uuid4().hex[:10],original).fetchone()[0]
    cursor.execute('INSERT INTO DetalleFacturas(FacturaID,ProductoID,Cantidad,PrecioUnitario,Subtotal) VALUES(?,?,1,7,7)',rid,product)
    document = json.loads(cursor.execute('SELECT DocumentoJSON FROM Facturas WHERE ID=?',case['facturas'][0]).fetchone()[0])
    document.update(id=rid,numero_factura=number,fecha=local.strftime('%d/%m/%Y'),hora=local.strftime('%H:%M:%S'),
                    asiento_id=entry['asiento_id'],subtotal=7,iva=0,total=7)
    document['productos'][0].update(precio_lista=7,subtotal=7,iva=0)
    cursor.execute('UPDATE Facturas SET DocumentoJSON=? WHERE ID=?',json.dumps(document),rid)
    case['conn'].commit()
    before = repo.snapshot(cursor,{'Desde':start,'Hasta':end},'Diagnóstico interno')
    record = next(r for r in before['ventas'] if r['id']==rid)
    assert record['fecha_normalizada'] and record['fecha'] == local.isoformat()
    assert record['fecha_bd'] == original.isoformat()
    next_start = end + timedelta(days=1)
    next_end = next_start.replace(day=calendar.monthrange(next_start.year,next_start.month)[1])
    after = repo.snapshot(cursor,{'Desde':next_start,'Hasta':next_end},'Diagnóstico interno')
    assert all(r['id']!=rid for r in after['ventas'])
    assert cursor.execute('SELECT FechaHora FROM Facturas WHERE ID=?',rid).fetchone()[0] == original


def test_new_invoice_and_ledger_share_nicaragua_timestamp_at_month_boundary(client,case,monkeypatch):
    from app.pos import routes as pos_routes
    cursor = case['conn'].cursor()
    _, end = repo.monthly_period(case['mes'])
    stamp = datetime(end.year,end.month,end.day,23,59,59)
    monkeypatch.setattr(pos_routes,'business_now',lambda:stamp)
    token = csrf(client)
    headers = {'X-CSRFToken':token}
    opened = client.post('/api/turno/abrir',json={},headers=headers)
    assert opened.status_code in (200,201),opened.json
    product = cursor.execute('SELECT TOP 1 ID,PrecioBase FROM Productos WHERE Activo=1 AND EsFicticio=0 AND PrecioBase>0 ORDER BY ID').fetchone()
    response = client.post('/api/factura',json={'clave_solicitud':uuid.uuid4().hex,
        'carrito':[{'id':product[0],'cantidad':1}],
        'pagos_multiples':[{'metodo':'Efectivo','monto':float(product[1])}]},headers=headers)
    assert response.status_code == 200,response.json
    document = response.json['factura']
    assert document['zona_horaria'] == 'America/Managua'
    assert document['fecha'] == stamp.strftime('%d/%m/%Y')
    assert cursor.execute('SELECT FechaHora FROM Facturas WHERE ID=?',document['id']).fetchone()[0] == stamp
    assert cursor.execute('SELECT Fecha FROM AsientosDiario WHERE ID=?',document['asiento_id']).fetchone()[0] == stamp


def test_migration_roundtrip_and_refusal_to_drop_nonempty_history(qa_string):
    # Base efímera separada: nunca se borra el esquema de la aplicación o de su QA principal.
    admin = pyodbc.connect(qa_string,autocommit=True)
    admin.execute('USE master')
    name = 'PanaderiaQA_DGIMigration_'+uuid.uuid4().hex[:10]
    admin.execute(f'CREATE DATABASE [{name}]')
    connection = pyodbc.connect(qa_string)
    try:
        connection.execute(f'USE [{name}]')
        require_qa(connection)
        connection.execute('''CREATE TABLE Usuarios(ID INT PRIMARY KEY);
            INSERT INTO Usuarios VALUES(1);
            CREATE TABLE ConfiguracionSistema(Clave NVARCHAR(100) PRIMARY KEY,Valor NVARCHAR(MAX),Descripcion NVARCHAR(255));
            CREATE TABLE Facturas(DocumentoJSON NVARCHAR(MAX));
            CREATE TABLE ComprasMateriaPrima(IVADocumentado DECIMAL(18,2));
            CREATE TABLE AsientosDiario(NumeroComprobante NVARCHAR(20));
            CREATE TABLE BitacoraFiscal(Detalle NVARCHAR(MAX));''')
        connection.commit()
        migrate(connection); migrate(connection); connection.commit()
        downgrade(connection); connection.commit()
        assert connection.execute("SELECT OBJECT_ID('DGI_Revisiones')").fetchone()[0] is None
        migrate(connection); connection.commit()
        connection.execute("INSERT INTO DGI_Periodos(Desde,Hasta,CreadoPor) VALUES('2026-10-01','2026-10-31',1)")
        connection.commit()
        with pytest.raises(ValueError,match='hay datos DGI'):
            downgrade(connection)
        connection.rollback()
        assert connection.execute('SELECT COUNT(*) FROM DGI_Periodos').fetchone()[0] == 1
    finally:
        connection.rollback()
        connection.autocommit = True
        connection.execute('SET TRANSACTION ISOLATION LEVEL READ COMMITTED')
        connection.execute('USE master')
        connection.close()
        admin.execute(f'DROP DATABASE [{name}]')
        admin.close()
