"""Casos de negocio independientes de la base configurada por el usuario."""
import copy
import io
import zipfile
from datetime import date
from decimal import Decimal
from xml.etree import ElementTree as ET

import pytest
from app.services.dgi_engine import analyze, fingerprint, json_dump, monthly_period, money
from app.services.dgi_excel import export_xlsx


@pytest.fixture
def snapshot():
    return dict(desde='2026-10-01', hasta='2026-10-31', obligacion='Diagnóstico interno',
        perfil=dict(fiscal_razon_social='Empresa sintética', fiscal_ruc='FICTICIO', fiscal_regimen='Cuota fija',
            dgi_responsable='Revisor QA', dgi_obligaciones='Diagnóstico interno', dgi_vigente_desde='2026-01-01', dgi_vigente_hasta=''),
        ventas=[dict(id=i, numero=f'DEMO-{i}', fecha='2026-10-10T10:00:00', estado='Emitida',
            anulada_en=None, base='40000.00', iva='0.00', total='40000.00', documento=True,
            detalles=[dict(subtotal='40000.00',cantidad=1,precio='40000.00')]) for i in (1, 2, 3)],
        compras=[dict(id=1, numero='C-DEMO', fecha='2026-10-10T10:00:00', proveedor_id=1, proveedor='Proveedor ficticio',
            materia_id=1, cantidad='1', total='5000', iva='0')],
        asientos=[dict(id=1, numero='FV-DEMO', fecha='2026-10-10T10:00:00', estado='Contabilizado', referencia='DEMO',
            detalles=[dict(cuenta='CAJA', nombre='Caja', clase='Activo', debe='110000', haber='0'),
                      dict(cuenta='VENTAS', nombre='Ventas', clase='Ingreso', debe='0', haber='110000')])])


def codes(result):
    return {f['codigo'] for f in result['hallazgos']}


def test_difference_is_10000_without_inventing_tax(snapshot):
    result = analyze(snapshot)
    assert Decimal(result['totales']['diferencia']) == 10000
    finding = next(f for f in result['hallazgos'] if f['codigo'] == 'CONCILIACION')
    assert finding['severidad'] == 'ERROR_BLOQUEANTE' and finding['diferencia'] == '10000.00'
    assert 'no es impuesto ni multa' in finding['accion']
    assert not any('tasa' in key or 'impuesto_pagar' in key for key in result['totales'])


def test_reconciliation_compares_base_excluding_recorded_tax(snapshot):
    for row in snapshot['ventas']:
        row.update(iva='6000', total='46000')
    snapshot['asientos'][0]['detalles'][0]['debe'] = '120000'
    snapshot['asientos'][0]['detalles'][1]['haber'] = '120000'
    result = analyze(snapshot)
    assert result['bloqueantes'] == 0
    assert Decimal(result['totales']['ventas_total']) == 138000
    assert Decimal(result['totales']['diferencia']) == 0


def test_empty_period_does_not_invent_liability(snapshot):
    snapshot.update(ventas=[], compras=[], asientos=[])
    result = analyze(snapshot)
    assert result['bloqueantes'] == 0
    assert all(Decimal(v) == 0 for v in result['totales'].values())
    assert 'SIN_COMPRAS' in codes(result)


def test_missing_profile_is_configuration_not_a_data_error(snapshot):
    snapshot['perfil'] = {}
    snapshot['ventas'][0]['base'] = snapshot['ventas'][0]['total'] = '30000.00'
    snapshot['ventas'][0]['detalles'][0].update(subtotal='30000.00',precio='30000.00')
    result = analyze(snapshot)
    assert 'PERFIL' in codes(result)
    assert result['bloqueantes'] == 0 and result['datos_consistentes']
    assert next(f for f in result['hallazgos'] if f['codigo']=='PERFIL')['categoria'] == 'CONFIGURACION'


def test_profile_must_cover_full_period(snapshot):
    snapshot['perfil']['dgi_vigente_hasta'] = '2026-10-15'
    assert 'VIGENCIA' in codes(analyze(snapshot))


def test_duplicate_invoice_references_source(snapshot):
    snapshot['ventas'][1]['numero'] = 'DEMO-1'
    finding = next(f for f in analyze(snapshot)['hallazgos'] if f['codigo'] == 'DUPLICADO')
    assert finding['origen_id'] == 2 and '#1' in finding['detalle']


def test_missing_number_and_purchase_support(snapshot):
    snapshot['ventas'][0]['numero'] = ''
    snapshot['compras'][0]['proveedor'] = ''
    assert {'NUMERO', 'SOPORTE'} <= codes(analyze(snapshot))


def test_arithmetic_and_details_are_separate_findings(snapshot):
    snapshot['ventas'][0]['total'] = '40001'
    snapshot['ventas'][0]['detalles'][0]['subtotal'] = '1'
    assert {'ARITMETICA', 'DETALLE'} <= codes(analyze(snapshot))


def test_out_of_period_sale_is_excluded(snapshot):
    snapshot['ventas'][0]['fecha'] = '2026-09-30T23:59:59'
    result = analyze(snapshot)
    assert Decimal(result['totales']['ventas_base']) == 80000
    assert not result['ventas'][0]['incluida']
    assert 'FUERA_PERIODO' in codes(result)


def test_same_period_void_and_reversal_cancel_without_double_counting(snapshot):
    row = snapshot['ventas'][0]
    row.update(estado='Anulada', anulada_en='2026-10-11T10:00:00')
    snapshot['asientos'][0]['estado'] = 'Anulado'
    snapshot['asientos'][0]['detalles'][0]['debe'] = '120000'
    snapshot['asientos'][0]['detalles'][1]['haber'] = '120000'
    reverse = copy.deepcopy(snapshot['asientos'][0])
    reverse.update(id=2, estado='Contabilizado')
    reverse['detalles'][0].update(debe='0', haber='40000')
    reverse['detalles'][1].update(debe='40000', haber='0')
    snapshot['asientos'].append(reverse)
    result = analyze(snapshot)
    assert result['bloqueantes'] == 0 and Decimal(result['totales']['diferencia']) == 0


@pytest.mark.parametrize('issued,voided', [('2026-10-10','2026-11-02'), ('2026-09-25','2026-10-10'), ('2026-10-10',None)])
def test_void_other_period_or_without_date_requires_review(snapshot, issued, voided):
    snapshot['ventas'][0].update(fecha=issued, estado='Anulada', anulada_en=voided)
    assert 'ANULACION' in codes(analyze(snapshot))


def test_purchase_lines_of_different_materials_are_not_duplicate(snapshot):
    row = copy.deepcopy(snapshot['compras'][0])
    row.update(id=2, materia_id=2)
    snapshot['compras'].append(row)
    assert 'DUPLICADO' not in codes(analyze(snapshot))
    row['materia_id'] = 1
    assert 'DUPLICADO' in codes(analyze(snapshot))


def test_unbalanced_journal_blocks(snapshot):
    snapshot['asientos'][0]['detalles'][0]['debe'] = '120000'
    assert 'PARTIDA_DOBLE' in codes(analyze(snapshot))


@pytest.mark.parametrize('amounts', [
    [('10', '10')],
    [('0', '0'), ('0', '0')],
    [('10', '10'), ('10', '10')],
    [('10', '0'), ('0', '10'), ('0', '0')],
    [('-10', '0'), ('0', '-10')],
    [],
])
def test_structurally_invalid_balanced_journal_blocks_even_with_pending_profile(snapshot, amounts):
    snapshot.update(perfil={}, ventas=[], compras=[])
    snapshot['asientos'][0]['detalles'] = [dict(cuenta=f'C-{i}', nombre='Sintética', clase='Activo',
        debe=debit, haber=credit) for i, (debit, credit) in enumerate(amounts)]
    before = copy.deepcopy(snapshot)
    result = analyze(snapshot)
    finding = next((f for f in result['hallazgos'] if f['codigo'] == 'PARTIDA_DOBLE'), None)
    assert finding is not None
    assert finding['severidad'] == 'ERROR_BLOQUEANTE' and finding['origen_id'] == 1
    assert not result['datos_consistentes']
    assert result['pendientes_configuracion'] == 1
    assert result == analyze(snapshot) and snapshot == before


def test_journal_validation_preserves_exact_decimal_precision(snapshot):
    snapshot.update(perfil={}, ventas=[], compras=[])
    snapshot['asientos'][0]['detalles'] = [
        dict(cuenta='A', nombre='Caja', clase='Activo', debe='0.00001', haber='0'),
        dict(cuenta='B', nombre='Pasivo', clase='Pasivo', debe='0', haber='0.00001')]
    assert analyze(snapshot)['datos_consistentes']
    snapshot['asientos'][0]['detalles'][1]['haber'] = '0.00002'
    assert 'PARTIDA_DOBLE' in codes(analyze(snapshot))


def test_review_from_engine_accepting_invalid_journals_must_be_regenerated():
    from app.services import dgi_repository as repository

    class EmptyPeriodCursor:
        def __init__(self):
            self.invalidated = False

        def execute(self, sql, *params):
            self.description, self.result = [], []
            if 'SELECT ID,Desde,Hasta,Estado FROM DGI_Periodos' in sql:
                self.description = [(name,) for name in ('ID', 'Desde', 'Hasta', 'Estado')]
                self.result = [(1, date(2026, 10, 1), date(2026, 10, 31), 'Abierto')]
            elif "UPDATE DGI_Revisiones SET Estado='OBSOLETO'" in sql:
                self.invalidated = True
            return self

        def fetchall(self):
            return self.result

    cursor = EmptyPeriodCursor()
    period = {'Desde': date(2026, 10, 1), 'Hasta': date(2026, 10, 31)}
    source = repository.snapshot(cursor, period, 'Diagnóstico interno')
    review = dict(ID=8, PeriodoID=1, Obligacion='Diagnóstico interno', Estado='EN_REVISION',
        Motor='dgi-interno-1.1', Huella=fingerprint(source))
    with pytest.raises(repository.DGIStale):
        repository.assert_fresh(cursor, review, 9)
    assert cursor.invalidated


def test_draft_journal_is_not_counted(snapshot):
    snapshot['asientos'][0]['estado'] = 'Borrador'
    result = analyze(snapshot)
    assert Decimal(result['totales']['ingresos_contables']) == 0
    assert {'BORRADOR_CONTABLE', 'SIN_CONTABILIDAD'} <= codes(result)


def test_deterministic_result_and_hash_without_mutating_source(snapshot):
    before = copy.deepcopy(snapshot)
    assert analyze(snapshot) == analyze(snapshot)
    assert snapshot == before
    assert fingerprint(snapshot) == fingerprint(dict(reversed(list(snapshot.items()))))
    snapshot['compras'][0]['iva'] = '1'
    assert fingerprint(snapshot) != fingerprint(before)


@pytest.mark.parametrize('value', ['NaN','Infinity','-Infinity',None,'bad','9999999999999999999'])
def test_invalid_money_is_not_silently_zero(value):
    with pytest.raises(ValueError):
        money(value)


def test_leap_month_and_invalid_month():
    assert monthly_period('2028-02') == (date(2028,2,1), date(2028,2,29))
    with pytest.raises(ValueError):
        monthly_period('2026-13')


def test_xlsx_is_readable_typed_and_has_no_formula_injection(snapshot):
    from openpyxl import load_workbook  # lector independiente, no autor del archivo
    snapshot['perfil']['fiscal_razon_social'] = ' \t=HYPERLINK("bad")'
    snapshot['compras'][0]['numero'] = '+cmd|test'
    snapshot['compras'][0]['proveedor'] = '@SUM(1,2)'
    result = analyze(snapshot)
    review = dict(ID=1,Version=1,Estado='DEMO SINTÉTICA',Huella=fingerprint(snapshot),CreadoPor=1)
    blob = export_xlsx(review, snapshot, result)
    book = load_workbook(io.BytesIO(blob), data_only=False)
    assert book.sheetnames == ['Resumen','Ventas','Compras','Hallazgos','Trazabilidad','Supuestos']
    assert book['Resumen']['B5'].value == ' \t=HYPERLINK("bad")'
    assert book['Ventas']['D5'].value == 40000
    assert book['Ventas']['C5'].value == datetime_for_test('2026-10-10T10:00:00')
    assert book['Compras']['B5'].value == '+cmd|test'
    assert book['Compras']['D5'].data_type == 's'
    assert 'NO VÁLIDA' in book['Resumen']['A1'].value
    assert all(c.data_type != 'f' for sheet in book for row in sheet for c in row)
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            if name.endswith('.xml'):
                ET.fromstring(archive.read(name))


def datetime_for_test(value):
    from datetime import datetime
    return datetime.fromisoformat(value)


def test_exact_user_case_46_sales_46_ledger_empty_profile_and_no_purchases(snapshot):
    snapshot['perfil'] = {}
    snapshot['compras'] = []
    for row, value in zip(snapshot['ventas'], ('36.00','5.00','5.00')):
        row.update(base=value,total=value)
        row['detalles'][0].update(precio=value,subtotal=value)
    snapshot['asientos'][0]['detalles'][0]['debe'] = '46'
    snapshot['asientos'][0]['detalles'][1]['haber'] = '46'
    result = analyze(snapshot)
    assert Decimal(result['totales']['diferencia']) == 0
    assert result['bloqueantes'] == result['advertencias'] == 0
    assert result['pendientes_configuracion'] == 1 and result['datos_consistentes']
    assert all(f['categoria'] != 'DATOS' for f in result['hallazgos'])


def test_date_without_evidence_and_corrupt_document_are_still_reported(snapshot):
    snapshot['ventas'][0].update(fecha_por_revisar=True,documento_corrupto=True)
    result = analyze(snapshot)
    assert {'FECHA_FUENTE','DOCUMENTO'} <= codes(result)
    assert any(f['codigo']=='DOCUMENTO' and f['severidad']=='ERROR_BLOQUEANTE' for f in result['hallazgos'])


def test_diagnostic_excel_keeps_errors_and_unapproved_label_even_invalid_amount(snapshot):
    from openpyxl import load_workbook
    snapshot['ventas'][0]['base'] = 'NaN'
    result = analyze(snapshot)
    review = dict(ID=1,Version=1,Estado='DIAGNÓSTICO SIN APROBACIÓN',Huella=fingerprint(snapshot),CreadoPor=1)
    book = load_workbook(io.BytesIO(export_xlsx(review,snapshot,result)))
    assert 'DIAGNÓSTICO SIN APROBACIÓN' in book['Ventas']['A2'].value
    assert book['Ventas']['D5'].value == 'DATO INVÁLIDO: NaN'
    assert any('IMPORTE' in str(c.value) for row in book['Hallazgos'] for c in row)
